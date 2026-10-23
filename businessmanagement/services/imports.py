import csv
import io
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction
from django.db.models import Q
from django.db.models.deletion import ProtectedError
from django.utils import timezone
from openpyxl import load_workbook

from tenantdata.models import (
    FinancialCategory,
    FinancialEntry,
    ImportBatch,
    ImportedRecord,
    Project,
    ThirdParty,
)

MAX_IMPORT_ROWS = 5000


def _header(value):
    value = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", value.strip().lower()).strip("_")


def _json_value(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


def read_import_rows(uploaded_file, delimiter=","):
    name = uploaded_file.name.lower()
    if name.endswith(".csv"):
        raw = uploaded_file.read()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
        reader = csv.DictReader(io.StringIO(text), delimiter=delimiter or ",")
        if not reader.fieldnames:
            raise ValidationError("El archivo debe tener una fila de encabezados.")
        rows = [dict(row) for row in reader]
    elif name.endswith(".xlsx"):
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = next(iterator, None)
        if not headers:
            raise ValidationError("El archivo Excel está vacío.")
        rows = [dict(zip(headers, values)) for values in iterator]
    else:
        raise ValidationError("Formato no soportado. Use CSV o Excel .xlsx.")

    normalized = []
    for row in rows:
        clean = {_header(key): _json_value(value) for key, value in row.items() if _header(key)}
        if any(str(value).strip() for value in clean.values()):
            normalized.append(clean)
    if not normalized:
        raise ValidationError("El archivo no contiene filas de datos.")
    if len(normalized) > MAX_IMPORT_ROWS:
        raise ValidationError(f"El archivo supera el máximo de {MAX_IMPORT_ROWS} filas.")
    return normalized


def _decimal(value):
    text = str(value or "").strip().replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    return Decimal(text)


def _date(value, required=False):
    if value in (None, ""):
        if required:
            raise ValueError
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    raise ValueError


def _boolean(value, default=True):
    if value in (None, ""):
        return default
    return str(value).strip().lower() in {"1", "true", "si", "sí", "yes", "activo"}


def validate_record(batch, row, seen_identifications=None):
    errors = []
    if batch.target_model == ImportBatch.THIRD_PARTY:
        name = str(row.get("name") or row.get("nombre") or "").strip()
        identification = str(row.get("identification") or row.get("identificacion") or "").strip()
        party_type = str(row.get("third_party_type") or row.get("tipo") or ThirdParty.CUSTOMER).strip().lower()
        email = str(row.get("email") or "").strip()
        if not name:
            errors.append("Falta la columna name/nombre.")
        if party_type not in dict(ThirdParty.TYPE_CHOICES):
            errors.append("third_party_type debe ser customer, supplier o both.")
        if email:
            try:
                validate_email(email)
            except ValidationError:
                errors.append("El email no es válido.")
        if identification:
            if ThirdParty.objects.filter(identification=identification).exists():
                errors.append("La identificación ya existe.")
            if seen_identifications is not None:
                if identification in seen_identifications:
                    errors.append("La identificación está repetida en el archivo.")
                seen_identifications.add(identification)
    else:
        concept = str(row.get("concept") or row.get("concepto") or "").strip()
        entry_type = str(row.get("entry_type") or row.get("tipo") or "").strip().lower()
        status = str(row.get("status") or FinancialEntry.PENDING).strip().lower()
        currency = str(row.get("currency") or row.get("moneda") or "COP").strip().upper()
        if not concept:
            errors.append("Falta concept/concepto.")
        try:
            if _decimal(row.get("amount", row.get("monto"))) <= 0:
                errors.append("El monto debe ser mayor que cero.")
        except (InvalidOperation, ValueError):
            errors.append("El monto no es válido.")
        if entry_type not in dict(FinancialEntry.ENTRY_TYPE_CHOICES):
            errors.append("entry_type no es válido.")
        if status not in dict(FinancialEntry.STATUS_CHOICES):
            errors.append("status no es válido.")
        if currency not in dict(FinancialEntry.CURRENCY_CHOICES):
            errors.append("currency no es válida.")
        try:
            entry_date = _date(row.get("entry_date", row.get("fecha")), required=True)
        except ValueError:
            entry_date = None
            errors.append("entry_date debe usar YYYY-MM-DD o DD/MM/YYYY.")
        try:
            due_date = _date(row.get("due_date", row.get("fecha_vencimiento")))
            if entry_date and due_date and due_date < entry_date:
                errors.append("due_date no puede ser anterior a entry_date.")
        except ValueError:
            errors.append("due_date no tiene un formato válido.")
        relations = (("category", "Categoría", FinancialCategory), ("project", "Proyecto", Project), ("third_party", "Tercero", ThirdParty))
        for key, label, model in relations:
            value = str(row.get(key) or "").strip()
            if value and not model.objects.filter(name__iexact=value).exists():
                errors.append(f"{label} '{value}' no existe en este tenant.")
    return errors


@transaction.atomic
def preview_batch(batch):
    batch = ImportBatch.objects.select_for_update().get(pk=batch.pk)
    if batch.status not in (ImportBatch.UPLOADED, ImportBatch.PREVIEWED):
        raise ValidationError("El lote ya no admite validación previa.")
    seen = set()
    valid = invalid = 0
    for record in batch.records.all():
        errors = validate_record(batch, record.row_data, seen)
        record.validation_errors = errors
        record.status = ImportedRecord.INVALID if errors else ImportedRecord.VALID
        record.save(update_fields=["validation_errors", "status"])
        if errors:
            invalid += 1
        else:
            valid += 1
    batch.status = ImportBatch.PREVIEWED
    batch.processed_rows = valid
    batch.failed_rows = invalid
    batch.save(update_fields=["status", "processed_rows", "failed_rows"])
    return batch


def _create_object(batch, row):
    if batch.target_model == ImportBatch.THIRD_PARTY:
        return ThirdParty.objects.create(
            name=str(row.get("name") or row.get("nombre")).strip(),
            identification=str(row.get("identification") or row.get("identificacion") or "").strip(),
            email=str(row.get("email") or "").strip(),
            phone=str(row.get("phone") or row.get("telefono") or "").strip(),
            address=str(row.get("address") or row.get("direccion") or "").strip(),
            third_party_type=str(row.get("third_party_type") or row.get("tipo") or ThirdParty.CUSTOMER).strip().lower(),
            is_active=_boolean(row.get("is_active"), True),
        )
    relation = {}
    for key, model in (("category", FinancialCategory), ("project", Project), ("third_party", ThirdParty)):
        value = str(row.get(key) or "").strip()
        if value:
            relation[key] = model.objects.get(name__iexact=value)
    return FinancialEntry.objects.create(
        business=batch.business,
        created_by=batch.created_by,
        entry_type=str(row.get("entry_type") or row.get("tipo")).strip().lower(),
        status=str(row.get("status") or FinancialEntry.PENDING).strip().lower(),
        concept=str(row.get("concept") or row.get("concepto")).strip(),
        amount=_decimal(row.get("amount", row.get("monto"))),
        currency=str(row.get("currency") or row.get("moneda") or "COP").strip().upper(),
        entry_date=_date(row.get("entry_date", row.get("fecha")), required=True),
        due_date=_date(row.get("due_date", row.get("fecha_vencimiento"))),
        notes=str(row.get("notes") or row.get("notas") or "").strip(),
        personal=_boolean(row.get("personal"), False),
        **relation,
    )


@transaction.atomic
def confirm_batch(batch):
    batch = ImportBatch.objects.select_for_update().get(pk=batch.pk)
    if batch.status != ImportBatch.PREVIEWED:
        raise ValidationError("El lote debe estar validado y sin confirmar.")
    records = list(batch.records.filter(status=ImportedRecord.VALID))
    if not records:
        raise ValidationError("No hay registros válidos para importar.")
    for record in records:
        obj = _create_object(batch, record.row_data)
        record.status = ImportedRecord.PROCESSED
        record.processed_object_id = str(obj.pk)
        record.processed_object_model = batch.target_model
        record.save(update_fields=["status", "processed_object_id", "processed_object_model"])
    batch.status = ImportBatch.CONFIRMED
    batch.processed_rows = len(records)
    batch.confirmed_at = timezone.now()
    batch.save(update_fields=["status", "processed_rows", "confirmed_at"])
    return batch


@transaction.atomic
def rollback_batch(batch, user):
    batch = ImportBatch.objects.select_for_update().get(pk=batch.pk)
    if batch.status != ImportBatch.CONFIRMED:
        raise ValidationError("Solo se puede revertir un lote confirmado.")
    model = FinancialEntry if batch.target_model == ImportBatch.FINANCIAL_ENTRY else ThirdParty
    records = list(batch.records.filter(status=ImportedRecord.PROCESSED))
    ids = [int(record.processed_object_id) for record in records if record.processed_object_id]
    existing_ids = set(model.objects.filter(pk__in=ids).values_list("pk", flat=True))
    if len(existing_ids) != len(ids):
        raise ValidationError("No se puede revertir: uno o más registros importados ya no existen.")
    if model is ThirdParty:
        used = ThirdParty.objects.filter(pk__in=ids).filter(
            Q(projects__isnull=False) | Q(financial_entries__isnull=False) |
            Q(invoices__isnull=False) | Q(quotes__isnull=False) |
            Q(credit_notes__isnull=False) | Q(debit_notes__isnull=False)
        ).distinct().exists()
        if used:
            raise ValidationError("No se puede revertir: un tercero importado ya está siendo utilizado.")
    try:
        model.objects.filter(pk__in=ids).delete()
    except ProtectedError as exc:
        raise ValidationError("No se puede revertir: un registro importado ya tiene operaciones relacionadas.") from exc
    batch.records.filter(status=ImportedRecord.PROCESSED).update(status=ImportedRecord.ROLLED_BACK)
    batch.status = ImportBatch.ROLLED_BACK
    batch.rolled_back_at = timezone.now()
    batch.rolled_back_by = user
    batch.save(update_fields=["status", "rolled_back_at", "rolled_back_by"])
    return batch
